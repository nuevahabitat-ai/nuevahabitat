"""Excel con URLs completas (links) de landings y blog para GSC."""
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SITE = 'https://www.nuevahabitat.com'
DESKTOP = Path.home() / 'Desktop'

LIKELY_NOT_INDEXED_LANDINGS = {
    'vender-horta', 'cuanto-vale-mi-piso-barcelona', 'nuevahabitat-vs-housfy-barcelona',
    'nuevahabitat-vs-clikalia-barcelona', 'vender-piso-herencia-barcelona',
    'vender-piso-rapido-barcelona', 'vender-piso-sin-exclusividad-barcelona',
    'nuevahabitat-vs-idealista-particular', 'vender-piso-alquilado-barcelona',
    'vender-el-clot-la-sagrera-barcelona', 'vender-cornella', 'vender-esplugues',
}

LIKELY_NOT_INDEXED_BLOG = {
    'gastos-compraventa', 'euribor-2026', 'home-staging', 'guia-hipotecas',
    'vender-piso-alquilado-guia-barcelona', 'vender-piso-divorcio-guia-barcelona',
    'idealista-fotocasa-vs-inmobiliaria-barcelona', 'vender-piso-poblenou-guia-2026',
    'guia-valoracion-piso-barcelona-2026', 'housfy-vs-precio-fijo-barcelona',
    'vender-piso-horta-guia', 'vender-hipoteca-pendiente-guia',
    'comision-inmobiliaria-barcelona', 'valoracion-gratis-barcelona',
}


def load_landings():
    index = json.loads((ROOT / 'content' / 'landings-index.json').read_text(encoding='utf-8'))
    landings_js = (ROOT / 'js' / 'landings.js').read_text(encoding='utf-8')
    m = re.search(r'window\.NH_LANDINGS = (\{[\s\S]*?\n\});', landings_js)
    meta = json.loads(m.group(1)) if m else {}
    rows = []
    for item in index.get('landings', []):
        slug = item['slug']
        cfg = meta.get(slug, {})
        url = f'{SITE}/{slug}'
        priority = slug in LIKELY_NOT_INDEXED_LANDINGS
        rows.append({
            'url': url,
            'titulo': cfg.get('footerLabel') or cfg.get('barrio') or slug,
            'tipo': 'Landing',
            'cluster': item.get('cluster', ''),
            'prioridad': 'ALTA' if priority else 'Normal',
        })
    return rows


def load_blog():
    src = (ROOT / 'js' / 'blog-posts.js').read_text(encoding='utf-8')
    rows = []
    for slug in re.findall(r"'([a-z0-9-]+)':\s*\{", src):
        title_m = re.search(rf"'{re.escape(slug)}':\s*\{{[\s\S]*?title:\s*'([^']*)'", src)
        cat_m = re.search(rf"'{re.escape(slug)}':\s*\{{[\s\S]*?cat:\s*'([^']*)'", src)
        priority = slug in LIKELY_NOT_INDEXED_BLOG
        rows.append({
            'url': f'{SITE}/blog/{slug}',
            'titulo': title_m.group(1) if title_m else slug,
            'tipo': 'Blog',
            'cluster': cat_m.group(1) if cat_m else '',
            'prioridad': 'ALTA' if priority else 'Normal',
        })
    return rows


def set_link_cell(cell, url):
    cell.value = url
    cell.hyperlink = url
    cell.font = Font(color='0563C1', underline='single')


def write_links_sheet(wb, title, rows, highlight_priority=False):
    ws = wb.create_sheet(title)
    headers = ['URL', 'Titulo', 'Tipo', 'Categoria', 'Prioridad indexacion']
    header_fill = PatternFill('solid', fgColor='1a1a1a')
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = header_fill
        c.font = Font(bold=True, color='FFFFFF')
    priority_fill = PatternFill('solid', fgColor='FFF3CD')
    for r, row in enumerate(rows, 2):
        set_link_cell(ws.cell(row=r, column=1), row['url'])
        ws.cell(row=r, column=2, value=row['titulo'])
        ws.cell(row=r, column=3, value=row['tipo'])
        ws.cell(row=r, column=4, value=row['cluster'])
        ws.cell(row=r, column=5, value=row['prioridad'])
        if highlight_priority and row['prioridad'] == 'ALTA':
            for col in range(1, 6):
                ws.cell(row=r, column=col).fill = priority_fill
    ws.column_dimensions['A'].width = 62
    ws.column_dimensions['B'].width = 55
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['E'].width = 18
    ws.freeze_panes = 'A2'
    return ws


def write_url_only_sheet(wb, title, urls):
    ws = wb.create_sheet(title)
    ws.cell(row=1, column=1, value='URL').font = Font(bold=True)
    ws.cell(row=1, column=2, value='Accion GSC').font = Font(bold=True)
    for r, url in enumerate(urls, 2):
        set_link_cell(ws.cell(row=r, column=1), url)
        ws.cell(row=r, column=2, value='Inspeccion URL > Solicitar indexacion')
    ws.column_dimensions['A'].width = 62
    ws.column_dimensions['B'].width = 38
    ws.freeze_panes = 'A2'


def main():
    landings = load_landings()
    blogs = load_blog()
    all_rows = landings + blogs
    priority_urls = [r['url'] for r in all_rows if r['prioridad'] == 'ALTA']

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    write_url_only_sheet(wb, '1-Prioridad GSC', priority_urls)
    write_links_sheet(wb, '2-Landings', landings, highlight_priority=True)
    write_links_sheet(wb, '3-Blog', blogs, highlight_priority=True)
    write_url_only_sheet(wb, '4-Todos landings', [r['url'] for r in landings])
    write_url_only_sheet(wb, '5-Todos blog', [r['url'] for r in blogs])

    outputs = [
        ROOT / 'imagenes' / 'NuevaHabitat-landings-blog-links.xlsx',
        DESKTOP / 'NuevaHabitat-landings-blog-links.xlsx',
    ]
    for out in outputs:
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out)
        print(f'Guardado: {out}')

    print(f'Landings: {len(landings)} URLs | Blog: {len(blogs)} URLs | Prioridad: {len(priority_urls)} URLs')


if __name__ == '__main__':
    main()
